# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# Palette
DARK   = RGBColor(0x1B, 0x2A, 0x4A)   # deep navy
ACCENT = RGBColor(0x5B, 0x8D, 0xEF)   # blue
ACC2   = RGBColor(0x7C, 0x5C, 0xE6)   # purple
LIGHT  = RGBColor(0xF3, 0xF5, 0xFA)
GRAY   = RGBColor(0x55, 0x5B, 0x66)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width  = Inches(13.333)
prs.slide_height = Inches(7.5)
SW, SH = prs.slide_width, prs.slide_height
BLANK = prs.slide_layouts[6]

def add_rect(slide, x, y, w, h, color, line=None):
    from pptx.enum.shapes import MSO_SHAPE
    sp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    sp.fill.solid(); sp.fill.fore_color.rgb = color
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line; sp.line.width = Pt(1)
    sp.shadow.inherit = False
    return sp

def txt(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(ln.get("sa", 6))
        p.space_before = Pt(ln.get("sb", 0))
        if ln.get("bullet"):
            p.level = ln.get("level", 0)
        for j, run in enumerate(ln["runs"]):
            r = p.add_run(); r.text = run[0]
            r.font.size = Pt(run[1]); r.font.bold = run[2]
            r.font.color.rgb = run[3]
            r.font.name = "Calibri"
    return tb

def bullets(slide, x, y, w, h, items, size=18, color=GRAY, gap=10):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        lvl = it[1] if isinstance(it, tuple) else 0
        text = it[0] if isinstance(it, tuple) else it
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap); p.level = lvl
        mark = "•  " if lvl == 0 else "–  "
        r = p.add_run(); r.text = mark + text
        r.font.size = Pt(size - lvl*2); r.font.color.rgb = color
        r.font.name = "Calibri"
    return tb

def header(slide, title, kicker=None):
    add_rect(slide, 0, 0, SW, Inches(1.15), DARK)
    add_rect(slide, 0, Inches(1.15), SW, Emu(48000), ACCENT)
    txt(slide, Inches(0.55), Inches(0.12), Inches(12), Inches(0.95),
        [{"runs":[(title, 26, True, WHITE)]}], anchor=MSO_ANCHOR.MIDDLE)
    if kicker:
        txt(slide, Inches(0.58), Inches(0.72), Inches(12), Inches(0.4),
            [{"runs":[(kicker, 12, False, RGBColor(0xB9,0xCB,0xEE))]}])

def footer(slide, n):
    txt(slide, Inches(0.5), Inches(7.05), Inches(9), Inches(0.35),
        [{"runs":[("Community-Level Blocklists in Decentralized Social Media — Zhang et al., 2025", 9, False, GRAY)]}])
    txt(slide, Inches(12.4), Inches(7.05), Inches(0.6), Inches(0.35),
        [{"runs":[(str(n), 10, True, ACCENT)]}], align=PP_ALIGN.RIGHT)

# ---------- SLIDE 1 : TITLE ----------
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, DARK)
add_rect(s, 0, Inches(4.55), SW, Emu(60000), ACCENT)
add_rect(s, 0, Inches(4.60), SW, Emu(30000), ACC2)
txt(s, Inches(0.9), Inches(1.6), Inches(11.5), Inches(2.6), [
    {"runs":[("Entendendo Blocklists de Nível Comunitário", 40, True, WHITE)], "sa":4},
    {"runs":[("em Redes Sociais Descentralizadas", 40, True, RGBColor(0x8F,0xB4,0xF2))], "sa":4},
])
txt(s, Inches(0.92), Inches(4.85), Inches(11.5), Inches(1.5), [
    {"runs":[("Zhang, Hwang, Liu, Horta Ribeiro & Monroy-Hernández — Princeton University (2025)", 16, False, RGBColor(0xC9,0xD6,0xF0))], "sa":6},
    {"runs":[("Moderação de conteúdo no Mastodon / Fediverse • arXiv:2506.05522", 13, False, GRAY)]},
])
txt(s, Inches(0.9), Inches(6.6), Inches(11), Inches(0.5),
    [{"runs":[("Apresentação de artigo", 12, True, ACCENT)]}])

# ---------- SLIDE 2 : CONTEXTO ----------
s = prs.slides.add_slide(BLANK); header(s, "Contexto e Motivação", "Por que estudar blocklists?")
bullets(s, Inches(0.6), Inches(1.5), Inches(7.1), Inches(5.2), [
    "Moderação = decidir o que aceitar ou rejeitar para manter o diálogo civil (o \"paradoxo da tolerância\" de Popper).",
    "Pesquisa anterior focou em plataformas centralizadas (X, Reddit) e em blocklists individuais.",
    "No Mastodon / Fediverse não existe \"plataforma\": milhares de instâncias independentes conectadas via protocolo ActivityPub.",
    "Cada instância define suas próprias regras — não há autoridade central de moderação.",
    "A principal ferramenta entre instâncias é a blocklist de nível comunitário: bloquear instâncias inteiras.",
], size=17, gap=13)
add_rect(s, Inches(8.05), Inches(1.6), Inches(4.65), Inches(4.9), LIGHT)
add_rect(s, Inches(8.05), Inches(1.6), Inches(0.12), Inches(4.9), ACC2)
txt(s, Inches(8.35), Inches(1.85), Inches(4.1), Inches(4.5), [
    {"runs":[("O impacto do bloqueio", 16, True, DARK)], "sa":10},
    {"runs":[("Bloquear instâncias grandes (ex.: mastodon.social) pode cortar a comunidade de grande parte do Fediverse.", 14, False, GRAY)], "sa":10},
    {"runs":[("Decisões de nível comunitário são muito mais impactantes que bloqueios individuais.", 14, False, GRAY)], "sa":10},
    {"runs":[("Falsos positivos têm consequências sérias.", 14, True, ACC2)]},
])
footer(s, 2)

# ---------- SLIDE 3 : RESEARCH QUESTIONS ----------
s = prs.slides.add_slide(BLANK); header(s, "Perguntas de Pesquisa", "Abordagem de métodos mistos")
rqs = [
    ("RQ1", "Qual é o panorama atual das blocklists de nível comunitário?", ACCENT),
    ("RQ2", "Como moderadores interpretam e aplicam as blocklists na prática?", ACC2),
    ("RQ3", "Que melhorias de design tornariam essas ferramentas mais úteis e eficazes?", RGBColor(0x2E,0xA0,0x7B)),
]
y = Inches(1.6)
for tag, q, col in rqs:
    add_rect(s, Inches(0.7), y, Inches(1.5), Inches(1.35), col)
    txt(s, Inches(0.7), y, Inches(1.5), Inches(1.35),
        [{"runs":[(tag, 24, True, WHITE)]}], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Inches(2.35), y, Inches(10.3), Inches(1.35), LIGHT)
    txt(s, Inches(2.65), y, Inches(9.8), Inches(1.35),
        [{"runs":[(q, 18, False, DARK)]}], anchor=MSO_ANCHOR.MIDDLE)
    y = Emu(y + Inches(1.6))
footer(s, 3)

# ---------- SLIDE 4 : MÉTODO ----------
s = prs.slides.add_slide(BLANK); header(s, "Metodologia", "Duas etapas complementares")
add_rect(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(4.9), LIGHT)
add_rect(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(0.7), ACCENT)
txt(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(0.7),
    [{"runs":[("1 · Análise de Conteúdo", 18, True, WHITE)]}], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(0.85), Inches(2.45), Inches(5.4), Inches(3.8), [
    "~8.700 instâncias Mastodon; analisadas 1.807 (≥10 usuários ativos).",
    "5 blocklists compartilhadas (Garden Fence, CARIAD, IFTAS-DNI, Seirdy Tier-0, FediNuke).",
    "4 ferramentas (FediCheck, FediBlockHole, Fediseer, The Bad Space).",
    "Categorias: propósito, critérios, transparência, distribuição, limitações.",
], size=15, gap=11)

add_rect(s, Inches(6.85), Inches(1.55), Inches(5.9), Inches(4.9), LIGHT)
add_rect(s, Inches(6.85), Inches(1.55), Inches(5.9), Inches(0.7), ACC2)
txt(s, Inches(6.85), Inches(1.55), Inches(5.9), Inches(0.7),
    [{"runs":[("2 · Entrevistas Semiestruturadas", 18, True, WHITE)]}], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
bullets(s, Inches(7.1), Inches(2.45), Inches(5.4), Inches(3.8), [
    "12 moderadores do Mastodon, 7 países, 26–55 anos.",
    "~1 hora cada, via Zoom; análise temática.",
    "1ª metade: práticas atuais de uso das blocklists (RQ2).",
    "2ª metade: protótipo (React/Vite) como \"design probe\" — filtros por categoria, níveis de severidade e comentários (RQ3).",
], size=15, gap=11)
footer(s, 4)

# ---------- SLIDE 5 : RQ1 landscape ----------
s = prs.slides.add_slide(BLANK); header(s, "RQ1 — Panorama das Blocklists", "Heterogêneas e pouco transparentes")
bullets(s, Inches(0.6), Inches(1.5), Inches(6.5), Inches(5.2), [
    "Só 20,1% das 1.807 instâncias compartilham publicamente sua blocklist.",
    "Entre essas, menos da metade (46,4%) dá o motivo do bloqueio → lacuna de transparência.",
    "Instâncias maiores compartilham mais (1000+ usuários: 45,3% vs. 10–24: 13,9%).",
    "Blocklists baseiam-se em curadoria manual e fontes externas, muitas vezes não divulgadas.",
    "Vieses recorrentes: foco no inglês, perspectiva do Norte Global, atualização lenta.",
], size=16, gap=12)
# top categories box
add_rect(s, Inches(7.4), Inches(1.55), Inches(5.35), Inches(4.95), LIGHT)
add_rect(s, Inches(7.4), Inches(1.55), Inches(5.35), Emu(45000), ACCENT)
cats = [("Spam","69,8%"),("Assédio / Troll","50,3%"),("Bots","47,3%"),
        ("Discurso de ódio","37,9%"),("CSAM / abuso infantil","33,7%"),
        ("Desinformação","29,6%"),("Transfobia","21,9%"),("Racismo","17,8%")]
lines = [{"runs":[("Top motivos de bloqueio", 16, True, DARK)], "sa":12}]
for name, pct in cats:
    lines.append({"runs":[(name+"  ", 14, False, GRAY), (pct, 14, True, ACC2)], "sa":7})
txt(s, Inches(7.7), Inches(1.85), Inches(4.8), Inches(4.5), lines)
footer(s, 5)

# ---------- SLIDE 6 : ferramentas ----------
s = prs.slides.add_slide(BLANK); header(s, "RQ1 — Blocklists e Ferramentas", "Diferentes filosofias de moderação")
rows = [
    ("Garden Fence / CARIAD", "Amplo consenso; domínios amplamente bloqueados (protege por abrangência)."),
    ("IFTAS-DNI / FediNuke", "Foco em severidade: instâncias notoriamente prejudiciais."),
    ("Seirdy Tier-0", "Consenso (≥15 de 27 listas); \"receipts\" documentando decisões."),
    ("FediCheck / FediBlockHole", "\"Moderação como serviço\": sincroniza listas confiáveis, usuário mantém controle."),
    ("Fediseer / The Bad Space", "Classificação e transparência; sinaliza instâncias via múltiplas comunidades."),
]
y = Inches(1.55)
for name, desc in rows:
    add_rect(s, Inches(0.6), y, Inches(3.8), Inches(0.95), DARK)
    txt(s, Inches(0.75), y, Inches(3.55), Inches(0.95),
        [{"runs":[(name, 14, True, WHITE)]}], anchor=MSO_ANCHOR.MIDDLE)
    add_rect(s, Inches(4.4), y, Inches(8.3), Inches(0.95), LIGHT)
    txt(s, Inches(4.6), y, Inches(7.95), Inches(0.95),
        [{"runs":[(desc, 14, False, GRAY)]}], anchor=MSO_ANCHOR.MIDDLE)
    y = Emu(y + Inches(1.03))
footer(s, 6)

# ---------- SLIDE 7 : RQ2 three approaches ----------
s = prs.slides.add_slide(BLANK); header(s, "RQ2 — Três Abordagens de Decisão", "Prioridades concorrentes na moderação")
cards = [
    ("Abertura", "Reativa, baseada em denúncias", "Valoriza conexão e livre expressão.", "Risco: exposição a atores nocivos.", ACCENT),
    ("Segurança", "Proativa, baseada em busca", "Bloqueio antecipado de más instâncias.", "Risco: penalizar atores de boa-fé.", ACC2),
    ("Contexto", "Revisão manual e cuidadosa", "Decisões justas e ponderadas.", "Risco: muito trabalhosa.", RGBColor(0x2E,0xA0,0x7B)),
]
x = Inches(0.6)
for title, sub, benefit, risk, col in cards:
    add_rect(s, x, Inches(1.65), Inches(3.95), Inches(4.7), LIGHT)
    add_rect(s, x, Inches(1.65), Inches(3.95), Inches(0.95), col)
    txt(s, x, Inches(1.65), Inches(3.95), Inches(0.95),
        [{"runs":[(title, 20, True, WHITE)]}], align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    txt(s, Emu(x+Inches(0.25)), Inches(2.8), Inches(3.45), Inches(3.4), [
        {"runs":[(sub, 14, True, DARK)], "sa":14},
        {"runs":[("✓ "+benefit, 14, False, GRAY)], "sa":12},
        {"runs":[("⚠ "+risk, 14, False, col)], "sa":6},
    ])
    x = Emu(x + Inches(4.12))
txt(s, Inches(0.6), Inches(6.5), Inches(12), Inches(0.6),
    [{"runs":[("Não são mutuamente exclusivas: moderadores alternam entre as abordagens conforme valores, metas e recursos da comunidade.", 14, True, DARK)]}])
footer(s, 7)

# ---------- SLIDE 8 : RQ3 needs ----------
s = prs.slides.add_slide(BLANK); header(s, "RQ3 — Necessidades de Design", "O que tornaria as blocklists melhores")
needs = [
    ("Colaboração comunitária", "Votação interna, listas de \"suspeitos\", bibliotecas públicas de \"receipts\", scorecards e um \"marketplace de blocklists\".", ACCENT),
    ("Eficiência e gestão", "Filtros por categoria (spam, assédio…), automação para detectar e aplicar bloqueios, integração ao Mastodon.", ACC2),
    ("Monitoramento e visibilidade", "Métricas por instância (atividade, população, denúncias), quem mais bloqueou, comentários e \"receipts\" com data/fonte.", RGBColor(0x2E,0xA0,0x7B)),
]
y = Inches(1.6)
for title, desc, col in needs:
    add_rect(s, Inches(0.6), y, Inches(0.18), Inches(1.5), col)
    add_rect(s, Inches(0.78), y, Inches(11.95), Inches(1.5), LIGHT)
    txt(s, Inches(1.05), Emu(y+Inches(0.12)), Inches(11.4), Inches(1.3), [
        {"runs":[(title, 17, True, DARK)], "sa":6},
        {"runs":[(desc, 15, False, GRAY)]},
    ])
    y = Emu(y + Inches(1.65))
footer(s, 8)

# ---------- SLIDE 9 : discussion ----------
s = prs.slides.add_slide(BLANK); header(s, "Discussão", "Implicações para moderação descentralizada")
bullets(s, Inches(0.6), Inches(1.5), Inches(12), Inches(5.2), [
    "Equilibrar prioridades concorrentes: abordagem híbrida integra abertura, segurança e contexto em vez de opô-las.",
    "Responsabilidade ampliada: bloqueios comunitários podem isolar comunidades inteiras — poder e risco maiores para moderadores.",
    ("O verdadeiro obstáculo não é falta de consciência dos moderadores, e sim a inadequação das ferramentas atuais.", 1),
    "Moderação colaborativa entre instâncias: documentação estruturada, tags universais, suporte multilíngue e \"covenants digitais\".",
    "Transparência confiável: métricas de domínio e metadados sobre origem e critérios das listas aumentam confiança e legitimidade.",
], size=17, gap=15)
footer(s, 9)

# ---------- SLIDE 10 : limitations + conclusion ----------
s = prs.slides.add_slide(BLANK); header(s, "Limitações e Conclusão")
add_rect(s, Inches(0.6), Inches(1.55), Inches(5.9), Inches(4.9), LIGHT)
txt(s, Inches(0.85), Inches(1.75), Inches(5.4), Inches(0.6),
    [{"runs":[("Limitações", 18, True, ACC2)]}])
bullets(s, Inches(0.85), Inches(2.4), Inches(5.4), Inches(3.9), [
    "Amostra pequena (12 moderadores) → generalização limitada.",
    "Foco em recursos populares e em inglês → possível viés de seleção.",
    "Práticas de moderação são dinâmicas e mudam ao longo do tempo.",
], size=15, gap=12)

add_rect(s, Inches(6.85), Inches(1.55), Inches(5.9), Inches(4.9), DARK)
txt(s, Inches(7.1), Inches(1.75), Inches(5.4), Inches(0.6),
    [{"runs":[("Conclusão", 18, True, RGBColor(0x8F,0xB4,0xF2))]}])
bullets(s, Inches(7.1), Inches(2.4), Inches(5.4), Inches(3.9), [
    "Blocklists comunitárias são essenciais e refletem filosofias diversas de moderação.",
    "Moderadores transitam entre abertura, segurança e contexto.",
    "Persistem desafios de esforço, transparência e colaboração.",
    "Futuro: ferramentas flexíveis, transparentes e colaborativas para comunidades descentralizadas resilientes.",
], size=15, color=RGBColor(0xD5,0xDE,0xF0), gap=12)
footer(s, 10)

# ---------- SLIDE 11 : end ----------
s = prs.slides.add_slide(BLANK)
add_rect(s, 0, 0, SW, SH, DARK)
add_rect(s, Inches(0.9), Inches(3.05), Inches(3.2), Emu(50000), ACCENT)
txt(s, Inches(0.9), Inches(2.4), Inches(11), Inches(1.2),
    [{"runs":[("Obrigado!", 44, True, WHITE)]}])
txt(s, Inches(0.92), Inches(3.4), Inches(11.5), Inches(2),
    [{"runs":[("Perguntas e discussão", 20, False, RGBColor(0xC9,0xD6,0xF0))], "sa":16},
     {"runs":[("Zhang, Hwang, Liu, Horta Ribeiro & Monroy-Hernández (2025) — Princeton University · arXiv:2506.05522", 13, False, GRAY)]}])

out = r"C:\Users\mlute\PycharmProjects\mestrado\PAPER2\apresentacao_blocklists.pptx"
prs.save(out)
print("Saved:", out)
